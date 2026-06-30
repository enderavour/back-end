from openpyxl import load_workbook
from fastapi import HTTPException

from app.schemas.quiz import (
    QuizCreate,
    QuestionCreate,
    AnswerOptionCreate,
)
from app.services.quiz import QuizService


class ImportService:

    @staticmethod
    async def import_excel(db, company_id, file):
        file.seek(0)

        workbook = load_workbook(file)
        sheet = workbook.active

        quizzes = ImportService._parse_sheet(sheet)

        if not quizzes:
            raise HTTPException(status_code=400, detail="Empty file")

        created = 0
        updated = 0

        for quiz_title, quiz_data in quizzes.items():

            existing = await QuizService.get_quiz_by_title(
                db,
                company_id,
                quiz_title
            )

            if existing:
                await QuizService.delete_quiz(db, existing)
                updated += 1
            else:
                created += 1

            quiz_schema = ImportService._build_quiz_schema(
                quiz_title,
                quiz_data
            )

            await QuizService.create_quiz(
                db,
                company_id,
                quiz_schema
            )

        return {
            "created": created,
            "updated": updated
        }

    @staticmethod
    def _parse_sheet(sheet):

        quizzes = {}

        for row in sheet.iter_rows(min_row=2, values_only=True):

            quiz_title = row[0]
            question_title = row[1]
            answer_title = row[2]
            is_correct = row[3]

            if quiz_title not in quizzes:
                quizzes[quiz_title] = {
                    "questions": {}
                }

            if question_title not in quizzes[quiz_title]["questions"]:
                quizzes[quiz_title]["questions"][question_title] = []

            quizzes[quiz_title]["questions"][question_title].append(
                {
                    "title": answer_title,
                    "is_correct": bool(is_correct)
                }
            )

        return quizzes

    @staticmethod
    def _build_quiz_schema(quiz_title, quiz_data):

        questions = []

        for question_title, answers in quiz_data["questions"].items():

            question = QuestionCreate(
                title=question_title,
                answers=[
                    AnswerOptionCreate(
                        title=answer["title"],
                        is_correct=answer["is_correct"]
                    )
                    for answer in answers
                ]
            )

            questions.append(question)

        return QuizCreate(
            title=quiz_title,
            description="Imported from Excel",
            questions=questions
        )
