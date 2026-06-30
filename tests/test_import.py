import pytest
from io import BytesIO
from openpyxl import Workbook
from app.services.importsrv import ImportService
from app.services.quiz import QuizService
from app.services.notification import NotificationService


def create_excel_file():
    wb = Workbook()
    ws = wb.active

    ws.append(["quiz", "question", "answer", "correct"])

    ws.append(["Python Quiz", "What is Python?", "Language", True])
    ws.append(["Python Quiz", "What is Python?", "Database", False])
    ws.append(["Python Quiz", "FastAPI is?", "Framework", True])
    ws.append(["Python Quiz", "FastAPI is?", "Browser", False])

    file = BytesIO()
    wb.save(file)
    file.seek(0)

    return file



@pytest.mark.asyncio
async def test_import_create_quiz(db_session):
    file = create_excel_file()

    result = await ImportService.import_excel(
        db_session,
        company_id=1,
        file=file
    )

    assert result["created"] >= 1


@pytest.mark.asyncio
async def test_import_update_existing_quiz(db_session):
    file = create_excel_file()

    await ImportService.import_excel(db_session, 1, file)

    file.seek(0)

    result = await ImportService.import_excel(db_session, 1, file)

    assert result["updated"] >= 1


@pytest.mark.asyncio
async def test_import_multiple_quizzes(db_session):
    wb = Workbook()
    ws = wb.active
    ws.append(["quiz", "question", "answer", "correct"])

    ws.append(["Quiz1", "Q1", "A1", True])
    ws.append(["Quiz1", "Q1", "A2", False])

    ws.append(["Quiz1", "Q2", "A3", True])
    ws.append(["Quiz1", "Q2", "A4", False])

    ws.append(["Quiz2", "Q1", "A1", False])
    ws.append(["Quiz2", "Q1", "A2", True])

    ws.append(["Quiz2", "Q2", "A3", False])
    ws.append(["Quiz2", "Q2", "A4", True])

    file = BytesIO()
    wb.save(file)
    file.seek(0)

    result = await ImportService.import_excel(db_session, 1, file)

    assert result["created"] >= 2


@pytest.mark.asyncio
async def test_import_empty_file(db_session):
    wb = Workbook()
    ws = wb.active
    ws.append(["quiz", "question", "answer", "correct"])

    file = BytesIO()
    wb.save(file)
    file.seek(0)

    with pytest.raises(Exception):
        await ImportService.import_excel(db_session, 1, file)
